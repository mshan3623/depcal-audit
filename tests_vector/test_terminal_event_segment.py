"""종료해 균등 재배분은 이벤트를 넘지 않는다 — 감사 2026-09-23 A3 (INC-14).

`settle_terminal_evenly`가 종료해 전체(12개월)를 균등화하면서, 그해 안의 부분양도·자본적지출
**이전** 달까지 이벤트 **이후** 기준 금액으로 바꿔 썼다. 연간 합계는 맞으므로 연도별 표와 거울
대조(core도 같은 동작)로는 안 보였고, 월별에서만 드러났다:
  · capex 종료해 7월: 1~6월에 증가분 상각이 당겨져 월 장부가가 **음수**(−1,499,500)
  · 부분양도 종료해 6월: 엑셀 시트1 처분 누계 5,219,609 vs 시트3 누계 하락 5,399,699 (180,090원)

규약: 미래의 사건은 지나간 달을 바꾸지 않는다(인과성). 균등 재배분은 종료해 중 **마지막 이벤트
이후 구간**에만 적용한다 — 잔재(정률 5%·정액 율×n≠1)는 이벤트 후 자산의 것이다. 이벤트가 종료해
밖(이전)에 있으면 종전과 같다.
"""
import pytest

from vcore.disposal import disposal_split
from vcore.monthly_schedule import monthly_events

ACQ = (2020, 1)
LIVES = [3, 5, 6, 7, 9]


def _terminal_months(cost, life, declining, fye):
    """자연완료 스케줄의 종료해 (연, 월) 목록."""
    rows = monthly_events(cost, life, *ACQ, fiscal_end_month=fye, declining=declining)
    last = rows[-1].fiscal_year
    return [(r.year, r.month) for r in rows if r.fiscal_year == last]


def _key(rows):
    return [(r.year, r.month, r.amount, r.acc, r.book) for r in rows]


def _cases():
    for life in LIVES:
        for declining in (False, True):
            for fye in (12, 6):
                yield life, declining, fye


@pytest.mark.parametrize("life,declining,fye", list(_cases()))
def test_partial_disposal_in_terminal_year_does_not_rewrite_earlier_months(life, declining, fye):
    """인과성: 양도월까지는 '그 달에 전부양도(절단)한 스케줄'과 같아야 한다."""
    cost, amount = 10_000_000, 6_000_000
    for (y, m) in _terminal_months(cost, life, declining, fye)[1:-1]:
        with_disp = monthly_events(cost, life, *ACQ, fye, declining, disp=(amount, y, m))
        cut = monthly_events(cost, life, *ACQ, fye, declining, disp=(None, y, m))
        assert _key(with_disp[:len(cut)]) == _key(cut), f"양도 {y}-{m}"


@pytest.mark.parametrize("life,declining,fye", list(_cases()))
def test_partial_disposal_split_equals_accumulated_drop(life, declining, fye):
    """엑셀 시트1(처분 제거 누계) == 시트3(양도 다음 달 누계 하락) — 같은 파일 안의 두 숫자."""
    cost, amount = 10_000_000, 6_000_000
    for (y, m) in _terminal_months(cost, life, declining, fye)[1:-1]:
        rows = monthly_events(cost, life, *ACQ, fye, declining, disp=(amount, y, m))
        i = next(k for k, r in enumerate(rows) if (r.year, r.month) == (y, m))
        drop = rows[i].acc + rows[i + 1].amount - rows[i + 1].acc
        _, want, _, _ = disposal_split(cost, rows[i].acc, rows[i].book, amount)
        assert drop == want, f"양도 {y}-{m}: 누계 하락 {drop:,} vs 분배 {want:,}"


@pytest.mark.parametrize("life,declining,fye", list(_cases()))
def test_capex_in_terminal_year_does_not_pull_depreciation_into_earlier_months(life, declining, fye):
    """증가 전 달은 증가를 모른다 — 장부가 음수·증가분 선상각 금지."""
    cost = 10_000_000
    for (y, m) in _terminal_months(cost, life, declining, fye)[1:-1]:
        rows = monthly_events(cost, life, *ACQ, fye, declining, inc=(cost // 2, y, m))
        i = next(k for k, r in enumerate(rows) if (r.year, r.month) == (y, m))
        pm = (y, m - 1) if m > 1 else (y - 1, 12)
        cut = monthly_events(cost, life, *ACQ, fye, declining, disp=(None, *pm))
        assert _key(rows[:i]) == _key(cut), f"증가 {y}-{m}"
        assert all(r.amount >= 0 and r.book >= 0 for r in rows), f"증가 {y}-{m}"


def test_audit_pin_capex_book_value_never_negative():
    """감사 재현: 종료해 7월 capex — 종전엔 12개월 전부 583,250, 6월 장부가 −1,499,500."""
    rows = monthly_events(10_000_000, 5, 2020, 1, inc=(5_000_000, 2024, 7))
    y24 = [r for r in rows if r.year == 2024]
    assert min(r.book for r in rows) >= 1_000
    # 증가 전: 원 자산 종료해 상각 (2,000,000 − 비망 1,000) // 12 — 증가분이 섞이지 않는다
    assert [r.amount for r in y24[:6]] == [166_583] * 6
    assert rows[-1].book == 1_000 and rows[-1].acc == 14_999_000


def test_audit_pin_excel_sheet1_matches_sheet3(tmp_path):
    """감사 재현: 종전 시트1 5,219,609 vs 시트3 누계 하락 5,399,699."""
    import openpyxl
    from asset_schedule_generator import generate_depreciation_schedule
    out = str(tmp_path / "s.xlsx")
    generate_depreciation_schedule("A", "2020-01-01", 10_000_000, 5, "유형자산", "정액법",
                                   "2024-06-30", 6_000_000, output_path=out)
    wb = openpyxl.load_workbook(out)
    s1 = {r[1]: r[2] for r in wb.worksheets[0].iter_rows(values_only=True) if r[1]}
    rows = [r for r in wb.worksheets[2].iter_rows(values_only=True)
            if r and r[0] == 2024 and isinstance(r[2], int)]
    jun, jul = rows[5], rows[6]
    drop = jun[3] + jul[2] - jul[3]
    assert s1["감가상각누계액 (처분부분)"] == drop


# ── core 오라클 거울 — 월 단위 ─────────────────────────────────────────────
# 연도별 거울은 연간 합계가 맞아 이 결함에 눈이 멀어 있었다. core도 종료해 전체를 균등화했기
# 때문에 월별 거울도 종료해 이벤트 표본이 없어 침묵했다. 같은 규약(마지막 이벤트 이후 구간만
# 균등)을 core가 독립 구현으로 따르는지 월 단위로 대조한다.
from core.depreciation_engine import calculate_depreciation_enhanced  # noqa: E402
from core.dep_common import AssetFinancials, AssetInfo, DepreciationMethod  # noqa: E402

_INFO = AssetInfo(asset_id="A", asset_name="자산", asset_category="비품")


def _core(cost, life, declining, fye, **ev):
    menum = DepreciationMethod.DECLINING_BALANCE if declining else DepreciationMethod.STRAIGHT_LINE
    fin = AssetFinancials(cost=cost, life_in_years=life, start_date="2020-01-15", method=menum, **ev)
    r = calculate_depreciation_enhanced(fin, _INFO, fiscal_year_end_month=fye)
    return [(m.year, m.month, m.numeric_depreciation, m.accumulated_depreciation, m.book_value)
            for m in r.schedule]


@pytest.mark.parametrize("life,declining,fye", list(_cases()))
def test_core_mirror_partial_disposal_in_terminal_year(life, declining, fye):
    cost, amount = 10_000_000, 6_000_000
    for (y, m) in _terminal_months(cost, life, declining, fye)[1:-1]:
        v = _key(monthly_events(cost, life, *ACQ, fye, declining, disp=(amount, y, m)))
        c = _core(cost, life, declining, fye, disposal_date=f"{y}-{m:02d}-01",
                  disposal_amount=amount)
        assert c == v, f"양도 {y}-{m}"


@pytest.mark.parametrize("life,declining,fye", list(_cases()))
def test_core_mirror_capex_in_terminal_year(life, declining, fye):
    cost = 10_000_000
    for (y, m) in _terminal_months(cost, life, declining, fye)[1:-1]:
        v = _key(monthly_events(cost, life, *ACQ, fye, declining, inc=(cost // 2, y, m)))
        c = _core(cost, life, declining, fye, increase_date=f"{y}-{m:02d}-01",
                  increase_amount=cost // 2)
        assert c == v, f"증가 {y}-{m}"
