"""검증 산출물의 출처·고지 — 감사조서 재현성 요건.

감사증거로 쓰이는 표는 "어느 엔진이·어느 파일을·언제" 계산했는지가 표 안에 남아야 하고,
도구의 지위(감사인 판단을 대체하지 않음)가 산출물에 동행해야 한다. 둘 다 없던 것을
2026-08-22에 신설했고, 이 파일이 다시 빠지는 것을 막는다.
"""
import os

import pytest
from openpyxl import load_workbook

from depverify import provenance
from depverify.report import DISCLAIMER, summary_text, write_xlsx
from depverify.verdict import verify_all
from vcore import __version__ as ENGINE_VERSION


def _asset(**over):
    a = dict(fy=2024, intang=False, cost=100_000_000, life=5, acq_y=2022, acq_m=3,
             disposed=False, exp_dep=20_000_000, exp_acc=56_666_666,
             exp_bk=43_333_334, prev_acc=36_666_666, asset_name="기계장치")
    a.update(over)
    return a


@pytest.fixture
def ledger(tmp_path):
    p = tmp_path / "대장.xlsx"
    p.write_bytes(b"ledger-bytes")
    return str(p)


# ── 출처 정보 ──────────────────────────────────────────────────────────────
def test_collect_records_engine_version(ledger):
    meta = provenance.collect(ledger, 2024, 12, 0)
    assert ENGINE_VERSION in meta["엔진 버전"]


def test_collect_records_file_digest_of_actual_bytes(ledger, tmp_path):
    """지문은 대장 내용에서 나온다 — 파일명이 같아도 내용이 다르면 달라야 한다."""
    import hashlib
    expected = hashlib.sha256(b"ledger-bytes").hexdigest()
    assert provenance.collect(ledger, 2024, 12, 0)["대장 SHA-256"] == expected

    with open(ledger, "wb") as fh:          # 같은 이름, 다른 내용
        fh.write(b"tampered")
    assert provenance.collect(ledger, 2024, 12, 0)["대장 SHA-256"] != expected


def test_collect_reports_unknown_revision_without_inventing(tmp_path, monkeypatch):
    """git이 없으면 '미상'을 말한다 — 없는 정보를 지어내지 않는다."""
    monkeypatch.setattr(provenance.subprocess, "run",
                        lambda *a, **k: (_ for _ in ()).throw(OSError("git 없음")))
    assert "미상" in provenance.source_revision()


def test_collect_records_tolerance_and_scope(ledger):
    meta = provenance.collect(ledger, 2024, 3, 500)
    assert "500" in meta["허용차 설정"] and "감사인" in meta["허용차 설정"]
    assert "별표 4" in meta["상각률 근거"]
    assert meta["결산월"] == "3월"


def test_missing_file_digest_is_unknown_not_crash(tmp_path):
    assert provenance.file_digest(str(tmp_path / "없는파일.xlsx")) == "미상"


# ── 보고서 반영 ────────────────────────────────────────────────────────────
def test_xlsx_has_disclaimer_sheet(tmp_path, ledger):
    out = str(tmp_path / "r.xlsx")
    write_xlsx(verify_all([_asset()]), 2024, out,
               provenance.collect(ledger, 2024, 12, 0))
    wb = load_workbook(out)
    assert "고지" in wb.sheetnames
    text = "\n".join(str(c[0].value) for c in wb["고지"].iter_rows() if c[0].value)
    assert "대체하지 않습니다" in text
    assert "검증불능" in text and "문제없음이 아니라" in text


def test_xlsx_summary_carries_provenance(tmp_path, ledger):
    out = str(tmp_path / "r.xlsx")
    meta = provenance.collect(ledger, 2024, 12, 0)
    write_xlsx(verify_all([_asset()]), 2024, out, meta)
    rows = {str(r[0].value): str(r[1].value)
            for r in load_workbook(out)["요약"].iter_rows(max_col=2) if r[0].value}
    assert rows["대장 SHA-256"] == meta["대장 SHA-256"]
    assert ENGINE_VERSION in rows["엔진 버전"]
    assert rows["실행 일시"] == meta["실행 일시"]


def test_report_without_meta_still_works(tmp_path):
    """meta는 선택 인자 — 없어도 보고서는 생성되고 고지는 남는다."""
    out = str(tmp_path / "r.xlsx")
    write_xlsx(verify_all([_asset()]), 2024, out)
    assert "고지" in load_workbook(out).sheetnames


def test_summary_text_always_carries_the_caveat():
    """meta가 없어도 '검증불능 ≠ 문제없음' 경고는 항상 붙는다."""
    text = summary_text(verify_all([_asset()]), 2024)
    assert "감사인의 판단을 대체하지 않습니다" in text


def test_summary_text_shows_engine_and_digest_when_meta_given(ledger):
    meta = provenance.collect(ledger, 2024, 12, 0)
    text = summary_text(verify_all([_asset()]), 2024, meta)
    assert ENGINE_VERSION in text
    assert meta["대장 SHA-256"][:16] in text


def test_disclaimer_names_the_out_of_scope_items():
    """범위 밖 항목이 고지에서 빠지면 사용자가 도구를 과신한다."""
    text = " ".join(DISCLAIMER)
    for item in ("세무조정", "업무용승용차", "중고자산", "사업연도", "K-IFRS"):
        assert item in text
