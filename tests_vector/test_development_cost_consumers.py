"""개발비(§26①6 1/n)를 웹·CLI·엑셀에서 고를 수 있게 (INC-20, 사용자 결정 2026-10-14).

종전엔 소비자 계층에 개발비 선택지가 없어 '무형자산'으로 넣으면 별표4 정액이 됐다(1억·3년:
33,300,000 vs 1/n 33,333,333). depverify만 1/n이라 같은 회사 명세서가 depverify에서 '차이'였다.
계산은 실행 경로 단일 진입점 `monthly_events(development_cost=True)` 하나로 한다.
자본적지출·부분양도는 개발비 재계산 규칙이 확립되지 않아 거절한다(depverify도 계산하지 않는다).
"""
import openpyxl
import pytest

from asset_schedule_generator import AssetInput, generate_depreciation_schedule
from vcore import intangible
from vcore.monthly_schedule import monthly_events


def _by_fy(cal):
    out = {}
    for m in cal:
        out[m.fiscal_year] = out.get(m.fiscal_year, 0) + m.amount
    return sorted(out.items())


@pytest.mark.parametrize("cost,life,m,fye", [(100_000_000, 3, 1, 12), (200_000_000, 7, 7, 12),
                                             (123_456_789, 20, 11, 3), (50_000_000, 1, 4, 6)])
def test_events_development_cost_equals_yearly_api(cost, life, m, fye):
    cal = monthly_events(cost, life, 2026, m, fye, development_cost=True)
    exp = [(r.fiscal_year, r.depreciation)
           for r in intangible.schedule_development_cost(cost, life, 2026, m, fye)]
    assert _by_fy(cal) == exp and cal[-1].book == 1000


def test_events_development_cost_full_disposal_equals_yearly_api():
    cal = monthly_events(100_000_000, 5, 2026, 1, 12, development_cost=True,
                         disp=(None, 2028, 6))
    exp = [(r.fiscal_year, r.depreciation) for r in
           intangible.schedule_full_disposal_development_cost(100_000_000, 5, 2026, 1, 2028, 6)]
    assert _by_fy(cal) == exp


@pytest.mark.parametrize("kw", [dict(inc=(10_000_000, 2027, 3)),
                                dict(disp=(40_000_000, 2027, 3)),
                                dict(declining=True)])
def test_events_development_cost_rejects_unestablished_combinations(kw):
    with pytest.raises(ValueError, match="개발비"):
        monthly_events(100_000_000, 5, 2026, 1, 12, development_cost=True, **kw)


def _sheets(path):
    wb = openpyxl.load_workbook(path)
    s1 = {r[1]: r[2] for r in wb.worksheets[0].iter_rows(values_only=True) if r[1]}
    s2 = {r[0]: r[2] for r in wb.worksheets[1].iter_rows(values_only=True) if isinstance(r[0], int)}
    return s1, s2


def test_excel_development_cost_uses_one_over_n(tmp_path):
    path = generate_depreciation_schedule('개발A', '2026-01-01', 100_000_000, 3, '개발비', '정액법',
                                          output_path=str(tmp_path / "d.xlsx"))
    s1, s2 = _sheets(path)
    assert s1["자산유형"] == "개발비"
    assert s2[2026] == 33_333_333                    # 무형자산(별표4)이면 33,300,000


@pytest.mark.parametrize("over,msg", [
    (dict(depreciation_method="정률법"), "정액"),
    (dict(useful_life=21), "1~20"),
    (dict(increase_date="2027-03-01", increase_amount=10_000_000), "자본적지출"),
    (dict(disposal_date="2027-03-31", disposal_amount=40_000_000), "부분"),
])
def test_asset_input_development_cost_guards(over, msg):
    kw = dict(asset_name='개발A', acquisition_date='2026-01-01', acquisition_cost=100_000_000,
              useful_life=3, asset_type='개발비', depreciation_method='정액법')
    kw.update(over)
    ok, message = AssetInput(**kw).validate()
    assert not ok and msg in message


def test_asset_input_development_cost_allows_one_year_life():
    ok, message = AssetInput(asset_name='개발A', acquisition_date='2026-01-01',
                             acquisition_cost=12_000_000, useful_life=1, asset_type='개발비',
                             depreciation_method='정액법').validate()
    assert ok, message


def test_web_accepts_development_cost(tmp_path, monkeypatch):
    import dep_cal_web
    monkeypatch.setattr(dep_cal_web, "OUTPUT_FOLDER", tmp_path)
    client = dep_cal_web.app.test_client()
    r = client.post("/calculate", json={
        "asset_name": "개발A", "acquisition_date": "2026-01-01", "acquisition_cost": 100_000_000,
        "useful_life": 3, "asset_type": "개발비", "depreciation_method": "정액법"})
    assert r.status_code == 200, r.get_json()


def test_cli_collects_development_cost_without_capex_or_partial_prompts(monkeypatch):
    """CLI: 개발비는 자산유형 → 내용연수(1~20) 순, 상각방법 정액 고정, 자본적지출·부분양도 질문 없음."""
    import depcal
    answers = iter(["개발A", "2026-01-01", "100,000,000",
                    "3",              # 4. 자산유형 = 개발비
                    "1",              # 5. 내용연수 1년 (개발비는 1년 허용)
                    "y", "2026-09-30",  # 처분 있음 — 처분유형은 묻지 않고 전체양도
                    ""])              # 결산월 기본
    monkeypatch.setattr("builtins.input", lambda _="": next(answers))
    info = depcal.collect_asset_info()
    assert (info["asset_type"], info["useful_life"], info["depreciation_method"]) == ("개발비", 1, "정액법")
    assert info["increase_date"] is None and info["disposal_amount"] is None
