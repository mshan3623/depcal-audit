"""depverify 완전성·서식 지문·거짓 경보 가드 — 로드맵 Phase 3 (감사 G7·G11·G19·G23·G24).

이 파일이 지키는 명제는 하나다: **모른다고 말할 수 있어야 한다.**
  · 대장에서 행이 사라져도 알아챌 채널이 있어야 한다 (G11 통제합계).
  · 컬럼이 없으면 없는 채로 굴리지 말고 무엇이 불가능해지는지 말해야 한다 (G23).
  · 셀을 해석할 수 없으면 추정하지 말아야 한다 (G24 숫자 날짜셀).
  · 비교식이 성립하지 않는 조합은 '차이'가 아니라 '검증불능'이다 (G7).
"""
import openpyxl
import pytest

from depverify.reader import read_ledger
from depverify.verdict import verify_asset

_COLUMNS = ["계정과목", "Code.1", "자산명", "취득일자", "기초가액", "전기말상각누계액",
            "신규취득및증가", "연수", "상감법", "당기상각비범위액", "당기말상각누계액",
            "당기말장부가액", "구분", "양도/폐기일", "전기말장부가액"]


def _asset_row(name, cost, prev_acc, dep, acc, bk, **over):
    row = {"계정과목": "비품", "Code.1": "1", "자산명": name, "취득일자": "2022-04-01",
           "기초가액": cost, "전기말상각누계액": prev_acc, "신규취득및증가": 0,
           "연수": 5, "상감법": "정액법", "당기상각비범위액": dep,
           "당기말상각누계액": acc, "당기말장부가액": bk, "구분": "미상각분",
           "양도/폐기일": None, "전기말장부가액": cost - prev_acc}
    row.update(over)
    return row


def _total_row(label, dep, acc, bk):
    """합계행 — 더존은 상감법을 비우고 자산명 칸에 '합  계'를 쓴다."""
    row = {c: None for c in _COLUMNS}
    row.update({"자산명": label, "당기상각비범위액": dep,
                "당기말상각누계액": acc, "당기말장부가액": bk})
    return row


def _write(tmp_path, rows, name="ledger.xlsx", extra_sheet=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "고정자산"
    ws.append(_COLUMNS)
    for r in rows:
        ws.append([r.get(c) for c in _COLUMNS])
    if extra_sheet:
        ws2 = wb.create_sheet("두번째")
        ws2.append(_COLUMNS)
        for r in rows:
            ws2.append([r.get(c) for c in _COLUMNS])
    path = tmp_path / name
    wb.save(path)
    return str(path)


# ── G11: 통제합계 대사 ─────────────────────────────────────────────────────

def test_control_total_matches_when_no_row_is_lost(tmp_path):
    """자산행 합 == 합계행이면 대사 통과 — 경고 없음."""
    rows = [_asset_row("자산A", 12_000_000, 4_200_000, 2_400_000, 6_600_000, 5_400_000),
            _asset_row("자산B", 6_000_000, 1_000_000, 1_200_000, 2_200_000, 3_800_000),
            _total_row("합  계", 3_600_000, 8_800_000, 9_200_000)]
    _, _, _, control = read_ledger(_write(tmp_path, rows), 2025)
    assert control.found and control.matched
    assert control.message() is None
    assert control.rows_summed == 2 and control.rows_unsummed == 0


def test_control_total_detects_dropped_asset_row(tmp_path):
    """자산행이 구조행으로 오판돼 사라져도 합계행이 그것을 드러낸다.

    '기계장치소계'는 상감법이 비어 있어 `_structural`이 구조행으로 본다. 자산행이었다면
    지금까지는 **어떤 검출 채널도 없었다** — 완전성은 기록을 뒤져서는 안 잡히기 때문이다.
    """
    rows = [_asset_row("자산A", 12_000_000, 4_200_000, 2_400_000, 6_600_000, 5_400_000),
            # 자산행이지만 상감법이 비어 구조행으로 오판된다
            _asset_row("기계장치소계", 6_000_000, 1_000_000, 1_200_000, 2_200_000, 3_800_000,
                       상감법=None),
            _total_row("합  계", 3_600_000, 8_800_000, 9_200_000)]
    _, _, structural, control = read_ledger(_write(tmp_path, rows), 2025)
    assert structural == 2                                  # 오판된 자산행 + 진짜 합계행
    assert control.found and not control.matched
    assert control.deltas == {"exp_dep": 1_200_000, "exp_acc": 2_200_000, "exp_bk": 3_800_000}
    msg = control.message()
    assert "대장 합계 대사 불일치" in msg and "Δ당기상각 +1,200,000" in msg


def test_control_total_absent_is_reported_as_unverified_not_matched(tmp_path):
    """합계행이 없으면 '일치'가 아니라 '대사 불가'다 — 없는 확신을 만들지 않는다."""
    rows = [_asset_row("자산A", 12_000_000, 4_200_000, 2_400_000, 6_600_000, 5_400_000)]
    _, _, _, control = read_ledger(_write(tmp_path, rows), 2025)
    assert not control.found and not control.matched
    assert "대사 불가" in control.message()


def test_subtotal_row_is_not_mistaken_for_grand_total(tmp_path):
    """계정과목별 '소계'는 통제합계가 아니다 — 소계를 총계로 쓰면 늘 불일치가 뜬다."""
    rows = [_asset_row("자산A", 12_000_000, 4_200_000, 2_400_000, 6_600_000, 5_400_000),
            _total_row("비품소계", 2_400_000, 6_600_000, 5_400_000),
            _asset_row("자산B", 6_000_000, 1_000_000, 1_200_000, 2_200_000, 3_800_000),
            _total_row("총  계", 3_600_000, 8_800_000, 9_200_000)]
    _, _, _, control = read_ledger(_write(tmp_path, rows), 2025)
    assert control.label == "총계" and control.matched


def test_unreadable_rows_still_count_toward_control_total(tmp_path):
    """검증불능 행도 합에 넣는다 — 빼고 더하면 합계가 안 맞는 게 당연해져 대사가 죽는다."""
    rows = [_asset_row("자산A", 12_000_000, 4_200_000, 2_400_000, 6_600_000, 5_400_000),
            _asset_row("폐기자산", 6_000_000, 1_000_000, 1_200_000, 2_200_000, 3_800_000,
                       구분="폐기"),
            _total_row("합  계", 3_600_000, 8_800_000, 9_200_000)]
    assets, unreadable, _, control = read_ledger(_write(tmp_path, rows), 2025)
    assert len(assets) == 1 and len(unreadable) == 1        # 폐기는 스코프외
    assert control.matched                                   # 그래도 합계는 맞는다
    assert control.rows_summed == 2


# ── G23: 서식 지문 ─────────────────────────────────────────────────────────

def test_missing_column_says_what_becomes_impossible(tmp_path):
    """필수 컬럼이 없으면 무엇이 불가능해지는지 말한다 — 종전엔 8개만 검사했다."""
    rows = [_asset_row("자산A", 12_000_000, 4_200_000, 2_400_000, 6_600_000, 5_400_000)]
    path = _write(tmp_path, rows)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.cell(row=1, column=_COLUMNS.index("구분") + 1, value="구분삭제됨")
    wb.save(path)
    with pytest.raises(ValueError) as e:
        read_ledger(path, 2025)
    assert "구분" in str(e.value) and "양도/보유 판별 불가" in str(e.value)


def test_unknown_mapping_key_is_rejected(tmp_path):
    """--mapping 오타 키는 조용히 무시되지 않는다 — 매핑을 줬는데 안 먹는 사고."""
    rows = [_asset_row("자산A", 12_000_000, 4_200_000, 2_400_000, 6_600_000, 5_400_000)]
    with pytest.raises(ValueError) as e:
        read_ledger(_write(tmp_path, rows), 2025, mapping={"prev_acc ": "전기말상각누계액"})
    assert "알 수 없는 키" in str(e.value)


def test_sheet_index_selects_second_sheet(tmp_path):
    """정수 시트 인덱스가 이름이 아니라 인덱스로 먹는다 (G19의 반대편 — 리더 계약)."""
    rows = [_asset_row("자산A", 12_000_000, 4_200_000, 2_400_000, 6_600_000, 5_400_000)]
    path = _write(tmp_path, rows, extra_sheet=True)
    assets, _, _, _ = read_ledger(path, 2025, sheet=1)
    assert len(assets) == 1


# ── G24: 숫자 날짜셀 ───────────────────────────────────────────────────────

@pytest.mark.parametrize("col,label", [("취득일자", "취득일자"), ("양도/폐기일", "양도/폐기일")])
def test_numeric_date_cell_is_rejected_not_epoch_parsed(tmp_path, col, label):
    """숫자 날짜셀은 1970-01-01로 조용히 둔갑하지 않고 사유를 남긴다.

    `pd.to_datetime(20220115)`는 나노초 epoch로 읽혀 1970이 된다. 판정은 '차이'로 뜨지만
    보고서에 취득 1970-01이 찍혀 원인을 가린다.
    """
    over = {col: 20_220_115}
    if col == "양도/폐기일":
        over["구분"] = "양도자산"
    rows = [_asset_row("자산A", 12_000_000, 4_200_000, 2_400_000, 6_600_000, 5_400_000, **over)]
    assets, unreadable, _, _ = read_ledger(_write(tmp_path, rows), 2025)
    assert assets == []
    _, reason = unreadable[0]
    assert "모호한날짜셀" in reason and label in reason


# ── G7: 비교식이 없는 조합은 '차이'가 아니라 '검증불능' ────────────────────

def _sold(**over):
    a = {"fy": 2024, "intang": False, "cost": 12_000_000, "life": 5,
         "acq_y": 2022, "acq_m": 1, "disposed": True, "disp_y": 2024, "disp_m": 6,
         "prev_acc": 4_800_000, "exp_dep": 1_200_000, "exp_acc": 6_000_000, "exp_bk": 0}
    a.update(over)
    return a


def test_intangible_disposal_is_unverifiable_not_a_difference():
    """무형 양도는 항상 '차이'였다 — 양도 환산식이 간접법(유형) 전제이기 때문.

    같은 수치가 유형이면 '일치'가 난다. 사유가 틀린 경보는 감사인 시간을 태운다.
    실측 앵커 0건이므로 비교식을 세울 수 없다고 정직하게 말한다.
    """
    v = verify_asset(_sold(intang=True, exp_acc=1_200_000))
    assert v.status == "검증불능"
    assert "비교식미확립" in v.reason and "무형 양도" in v.reason


def test_tangible_disposal_still_verifies():
    """유형 양도(상각중)는 종전대로 판정된다 — G7 수정이 실측 경로를 건드리지 않았다."""
    v = verify_asset(_sold())
    assert v.status == "일치", (v.status, v.d_dep, v.d_acc, v.d_bk)


def test_disposal_after_full_depreciation_is_unverifiable():
    """상각완료 후 양도도 항상 '차이'였다 — `d_bk = cost − exp_bk`는 장부가액이
    취득원가여야 일치라는 뜻이라, 자산이 제거돼 장부 0인 대장은 무엇을 적든 차이다."""
    v = verify_asset(_sold(fy=2030, disp_y=2030, exp_dep=0, exp_acc=0, exp_bk=0))
    assert v.status == "검증불능"
    assert "비교식미확립" in v.reason and "상각완료" in v.reason
