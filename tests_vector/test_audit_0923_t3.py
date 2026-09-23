"""감사 2026-09-23 트랜치 T3 — 국소 결함 4건의 회귀 가드 (M1·M3·M4·M5)."""
import openpyxl
import pytest

from vcore import straight_line
from vcore.separate_asset import schedule_separate_asset


# ── M1: 분리자산 경로에 취득연도 당기를 넣으면 거부 ────────────────────────
# 분리자산 경로(전기말 누계 → 당기 1년)는 이월자산용이다. 취득연도 당기를 넣으면 취득월 전
# 달까지 12개월로 셌다(7월 취득 → 정액 2,400,000, 정답 1,200,000). 운영 호출처는 0곳이고
# 당기 취득은 일반 경로가 처음부터 취득월부터 계산하므로, 값을 맞추지 않고 입력을 거부한다.

@pytest.mark.parametrize("declining", [False, True])
@pytest.mark.parametrize("acq_month", range(2, 13))
def test_separate_asset_rejects_acquisition_year_before_acquisition_month(declining, acq_month):
    with pytest.raises(ValueError, match="이월자산용"):
        schedule_separate_asset(12_000_000, 5, 2025, acq_month, 0, 2025, declining=declining)


def test_separate_asset_allows_january_acquisition_year_and_later_start():
    """1월 취득(12개월 = 정답)과 취득월 이후 개시(합병 승계)는 거부 대상이 아니다."""
    jan = schedule_separate_asset(12_000_000, 5, 2025, 1, 0, 2025)
    assert jan[0].months == 12 and jan[0].depreciation == straight_line.schedule(
        12_000_000, 5, 2025, 1)[0].depreciation
    later = schedule_separate_asset(12_000_000, 5, 2025, 3, 0, 2025, start_month=9)
    assert later[0].months == 4


# ── M3: depverify — 취득일 공란 한 행이 대장 전체를 멈추면 안 된다 ──────────
# 종전: `_date_cell`이 공백에 NaT를 돌려주고, try 밖의 `int(dt.year)`가 터져 exit 2
# ("cannot convert float NaN to integer", 행 번호 없음) — 나머지 자산이 하나도 검증되지 않았다.

def test_blank_acquisition_date_is_one_unreadable_row_not_a_crash(tmp_path, capsys):
    from depverify.__main__ import main
    from test_depverify_completeness import _matching_asset, _write
    rows = [_matching_asset("자산A"), _matching_asset("자산B")]
    rows[1]["취득일자"] = None
    out = str(tmp_path / "r.xlsx")
    code = main([_write(tmp_path, rows), "--fy", "2025", "--out", out])
    stdout = capsys.readouterr().out
    assert code == 1                                     # 검증불능이 있으면 1 (중단 2가 아님)
    assert "일치 1건" in stdout and "검증불능 1건" in stdout
    assert "자산B" in stdout and "취득일자 공백" in stdout


# ── M4: disposal_amount=0 — 시트1과 시트3이 같은 판단을 본다 ────────────────
# 종전: 시트3은 `disposal_amount or None`으로 0을 전부양도로, 시트1은 원값 0을 넘겨 부분양도
# (제거액 0/0/0)로 해석 — 한 파일 안에서 두 판단(판단 2벌 3회차). 정규화를 입력 한 곳으로.

def test_zero_disposal_amount_is_full_disposal_on_every_sheet(tmp_path):
    from asset_schedule_generator import generate_depreciation_schedule
    out = str(tmp_path / "z.xlsx")
    generate_depreciation_schedule("A", "2020-01-01", 10_000_000, 5, "유형자산", "정액법",
                                   "2024-06-30", 0, output_path=out)
    wb = openpyxl.load_workbook(out)
    s1 = {r[1]: r[2] for r in wb.worksheets[0].iter_rows(values_only=True) if r[1]}
    last = [r for r in wb.worksheets[2].iter_rows(values_only=True)
            if r and isinstance(r[0], int) and isinstance(r[2], int)][-1]
    assert s1["처분유형"] == "전체처분"
    assert s1["취득원가 (전체)"] == 10_000_000
    assert (last[0], last[1]) == (2024, "6월")                # 시트3도 양도월에서 끝난다


@pytest.mark.parametrize("cost,ratio", [(5_000, 0.01), (99, 1.0)])
def test_cli_ratio_that_rounds_to_zero_is_rejected(cost, ratio):
    """0원이 되는 비율은 '부분양도'라고 표시한 채 전부양도로 흘러가면 안 된다."""
    from depcal import partial_disposal_amount_from_ratio
    with pytest.raises(ValueError, match="0원"):
        partial_disposal_amount_from_ratio(cost, ratio)


def test_cli_ratio_normal_case():
    from depcal import partial_disposal_amount_from_ratio
    assert partial_disposal_amount_from_ratio(10_000_000, 60) == 6_000_000


# ── M5: 엑셀 수식 주입 — 사용자·대장 문자열이 수식으로 저장되면 안 된다 ────────
# 자산명 `=HYPERLINK(...)`이 수식 셀(data_type 'f')로 기록됐다. 명세서도, 고객 대장에서 자산명을
# 옮겨 적는 depverify 보고서(감사조서)도 같은 경로다. 두 산출물 모두 의도한 수식은 없다.
INJECT = '=HYPERLINK("http://evil.example","x")'


def _formula_cells(path):
    wb = openpyxl.load_workbook(path)
    return [(ws.title, c.coordinate) for ws in wb.worksheets
            for row in ws.iter_rows() for c in row if c.data_type == "f"]


def test_schedule_generator_writes_formula_like_names_as_text(tmp_path):
    from asset_schedule_generator import generate_depreciation_schedule
    out = str(tmp_path / "f.xlsx")
    generate_depreciation_schedule(INJECT, "2020-01-01", 10_000_000, 5, "유형자산", "정액법",
                                   output_path=out)
    assert _formula_cells(out) == []
    ws = openpyxl.load_workbook(out).worksheets[0]
    assert INJECT in [c.value for row in ws.iter_rows() for c in row]   # 값은 그대로 보존


def test_depverify_report_writes_ledger_names_as_text(tmp_path, capsys):
    from depverify.__main__ import main
    from test_depverify_completeness import _matching_asset, _write
    ledger = _write(tmp_path, [_matching_asset(INJECT)])
    wb = openpyxl.load_workbook(ledger)                  # 대장엔 '텍스트'로 들어 있어야 실제 경로다 —
    for c in wb.active[2]:                               # 수식 셀이면 pandas가 빈칸으로 읽어 공허 통과
        if c.value == INJECT:
            c.data_type = "s"
    wb.save(ledger)
    out = str(tmp_path / "r.xlsx")
    main([ledger, "--fy", "2025", "--out", out])
    capsys.readouterr()
    names = [c.value for row in openpyxl.load_workbook(out)["자산별"].iter_rows() for c in row]
    assert INJECT in names                               # 자산명이 실제로 보고서에 도달했다
    assert _formula_cells(out) == []
