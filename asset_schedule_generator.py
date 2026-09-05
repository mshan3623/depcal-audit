# -*- coding: utf-8 -*-
# Copyright 2026 Han Myeong Su
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""
자산 감가상각 월별 스케줄 자동 생성기
Single Asset Depreciation Schedule Excel Generator

입력받은 단일 자산 정보를 바탕으로 취득일부터 감가상각 종료일까지
월별 감가상각 스케줄을 생성하여 Excel 파일로 출력합니다.
"""

import os
import re
from datetime import datetime
from typing import Optional, List, Dict, Any
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.page import PageMargins


def _format_calc_method(raw_method: str) -> str:
    """계산방법 문자열을 가독성 좋게 변환"""
    method_map = {
        "한국세법정액법_취득월할계산": "정액법 (취득월할)",
        "한국세법정액법": "정액법",
        "한국세법정률법": "정률법",
        "한국세법정률법_취득월할계산": "정률법 (취득월할)",
        "한국세법무형자산정액법": "무형자산 정액법",
        "한국세법무형자산정액법_취득월할계산": "무형자산 정액법 (취득월할)",
        "정액법_최종월조정": "정액법 (최종월)",
        "정률법_최종월조정": "정률법 (최종월)",
        "정액법_자본적지출반영": "정액법 (자본적지출)",
        "정률법_자본적지출반영": "정률법 (자본적지출)",
        "정액법_부분양도": "정액법 (부분양도)",
        "정률법_부분양도": "정률법 (부분양도)",
    }
    # 매핑에 있으면 변환, 없으면 원본에서 불필요한 접두사 제거
    if raw_method in method_map:
        return method_map[raw_method]
    # 기본 정리: "한국세법" 접두사 제거 및 언더스코어를 괄호로
    cleaned = raw_method.replace("한국세법", "").replace("_", " (")
    if "(" in cleaned and not cleaned.endswith(")"):
        cleaned += ")"
    return cleaned.strip()

from vcore import disposal
from vcore.monthly_schedule import monthly_events


class AssetInput:
    """자산 입력 정보 클래스"""

    def __init__(
        self,
        asset_name: str,
        acquisition_date: str,  # YYYY-MM-DD
        acquisition_cost: int,
        useful_life: int,
        asset_type: str,  # "유형자산" or "무형자산"
        depreciation_method: str,  # "정액법" or "정률법"
        disposal_date: Optional[str] = None,  # YYYY-MM-DD
        disposal_amount: Optional[int] = None,
        increase_date: Optional[str] = None,  # YYYY-MM-DD
        increase_amount: Optional[int] = None,
        fiscal_year_end_month: int = 12  # 회사 결산월 (application 레이어 default=12, 한국 실무)
    ):
        self.asset_name = asset_name
        self.acquisition_date = acquisition_date
        self.acquisition_cost = acquisition_cost
        self.useful_life = useful_life
        self.asset_type = asset_type
        self.depreciation_method = depreciation_method
        self.disposal_date = disposal_date
        self.disposal_amount = disposal_amount
        self.increase_date = increase_date
        self.increase_amount = increase_amount
        self.fiscal_year_end_month = fiscal_year_end_month

    def validate(self) -> tuple[bool, str]:
        """입력 데이터 유효성 검증"""
        # 자산명 검증
        if not self.asset_name or not self.asset_name.strip():
            return False, "자산명을 입력해주세요"

        # 취득일자 검증
        try:
            datetime.strptime(self.acquisition_date, "%Y-%m-%d")
        except ValueError:
            return False, "취득일자는 YYYY-MM-DD 형식이어야 합니다"

        # 취득원가 검증
        if self.acquisition_cost <= 0:
            return False, "취득원가는 0보다 커야 합니다"

        # 내용연수 검증
        if not (2 <= self.useful_life <= 60):
            return False, "내용연수는 2~60년 사이여야 합니다"

        # 자산유형 검증
        if self.asset_type not in ["유형자산", "무형자산"]:
            return False, "자산유형은 '유형자산' 또는 '무형자산'이어야 합니다"

        # 감가상각방법 검증
        if self.depreciation_method not in ["정액법", "정률법"]:
            return False, "감가상각방법은 '정액법' 또는 '정률법'이어야 합니다"

        # 무형자산은 정액법만 가능
        if self.asset_type == "무형자산" and self.depreciation_method == "정률법":
            return False, "무형자산은 정액법만 적용 가능합니다"

        # 처분일자 검증 (선택사항)
        if self.disposal_date:
            try:
                disposal_dt = datetime.strptime(self.disposal_date, "%Y-%m-%d")
                acquisition_dt = datetime.strptime(self.acquisition_date, "%Y-%m-%d")
                if disposal_dt <= acquisition_dt:
                    return False, "처분일자는 취득일자보다 이후여야 합니다"
            except ValueError:
                return False, "처분일자는 YYYY-MM-DD 형식이어야 합니다"

        # 처분금액 검증 (선택사항)
        if self.disposal_amount is not None and self.disposal_amount < 0:
            return False, "처분금액은 0 이상이어야 합니다"

        # 증가일자 검증 (선택사항)
        if self.increase_date:
            try:
                increase_dt = datetime.strptime(self.increase_date, "%Y-%m-%d")
                acquisition_dt = datetime.strptime(self.acquisition_date, "%Y-%m-%d")
                if increase_dt <= acquisition_dt:
                    return False, "증가일자는 취득일자보다 이후여야 합니다"
            except ValueError:
                return False, "증가일자는 YYYY-MM-DD 형식이어야 합니다"

        # 증가금액 검증 (선택사항)
        if self.increase_amount is not None and self.increase_amount < 0:
            return False, "증가금액은 0 이상이어야 합니다"

        # 처분 후 자본적지출은 성립하지 않는다 — 두 진입점(웹·CLI)이 같이 막히도록 여기 한 곳에 둔다
        if self.disposal_date and self.increase_date:
            try:
                if (datetime.strptime(self.increase_date, "%Y-%m-%d")
                        > datetime.strptime(self.disposal_date, "%Y-%m-%d")):
                    return False, "증가일자(자본적지출)는 처분일자보다 이후일 수 없습니다"
            except ValueError:
                pass          # 형식 오류는 위의 개별 검증이 이미 잡는다

        # 자본적지출이 있으면 처분금액의 기준은 통합 취득원가다. 취득원가 이상 ~ 통합원가 미만
        # 구간은 '전부'인지 '일부'인지 입력만으로 갈리지 않아(감사 G4) 조용히 부분양도로
        # 해석되면 양도 후에도 상각이 계속된다 — 뜻을 되묻는다.
        if (self.disposal_date and self.disposal_amount is not None
                and self.increase_amount and self.increase_amount > 0):
            basis = self.acquisition_cost + self.increase_amount
            if self.acquisition_cost <= self.disposal_amount < basis:
                return False, (
                    f"자본적지출이 있어 처분금액은 통합 취득원가({basis:,}원) 기준입니다. "
                    f"전부양도면 처분금액을 비우거나 {basis:,}원 이상으로, "
                    f"부분양도면 {self.acquisition_cost:,}원 미만으로 지정하세요")

        # 결산월 검증
        if not isinstance(self.fiscal_year_end_month, int) or not (1 <= self.fiscal_year_end_month <= 12):
            return False, "결산월(fiscal_year_end_month)은 1~12 사이 정수여야 합니다"

        return True, "OK"

class MonthlyScheduleGenerator:
    """월별 감가상각 스케줄 생성기"""

    def __init__(self, asset_input: AssetInput):
        self.asset_input = asset_input

    def generate_schedule(self) -> List[Dict[str, Any]]:
        """
        월별 감가상각 스케줄 생성

        Returns:
            List of dict with keys: year, month, monthly_dep, accumulated_dep, book_value, calc_method
        """
        # 입력 검증
        is_valid, message = self.asset_input.validate()
        if not is_valid:
            raise ValueError(f"입력 데이터 오류: {message}")

        # 감가상각 계산 실행 — vcore 슬림 코어 (회계기간 명시, 양도월 포함=더존식)
        ai = self.asset_input
        acq = datetime.strptime(ai.acquisition_date, "%Y-%m-%d")
        declining = ai.depreciation_method == "정률법"

        inc = None
        if ai.increase_date and ai.increase_amount:
            d = datetime.strptime(ai.increase_date, "%Y-%m-%d")
            inc = (ai.increase_amount, d.year, d.month)
        disp = None
        if ai.disposal_date:
            d = datetime.strptime(ai.disposal_date, "%Y-%m-%d")
            disp = (ai.disposal_amount or None, d.year, d.month)   # 0/None = 전부양도

        cal = monthly_events(ai.acquisition_cost, ai.useful_life, acq.year, acq.month,
                             ai.fiscal_year_end_month, declining, inc=inc, disp=disp)
        if not cal:
            raise ValueError("감가상각 스케줄을 생성할 수 없습니다")

        if ai.asset_type == "무형자산":
            method_str = "한국세법무형자산정액법"
        else:
            method_str = "한국세법정률법" if declining else "한국세법정액법"

        # 월별 스케줄 변환
        return [{
            "year": m.year,
            "month": m.month,
            "fiscal_year": m.fiscal_year,   # 결산월 반영 회계연도 — 연도별 집계는 이 값으로 묶는다
            "monthly_dep": m.amount,
            "accumulated_dep": m.acc,
            "book_value": m.book,
            "calc_method": method_str,
            "note": ""
        } for m in cal]


class ExcelGenerator:
    """Excel 출력 생성기 - 3시트 구조"""

    def __init__(self, asset_input: AssetInput, schedule: List[Dict[str, Any]]):
        self.asset_input = asset_input
        self.schedule = schedule
        self.disposal_info = self._calculate_disposal_info()

    def generate_excel(self, output_path: str) -> str:
        """
        Excel 파일 생성 (3시트 구조)
        - Sheet 1: 입력 정보
        - Sheet 2: 연도별 집계표
        - Sheet 3: 월별 상세표

        Args:
            output_path: 출력 파일 경로 (없으면 자동 생성)

        Returns:
            생성된 파일 경로
        """
        wb = openpyxl.Workbook()

        # 기본 시트 제거 후 3개 시트 생성
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])

        ws1 = wb.create_sheet("1. 입력정보", 0)
        ws2 = wb.create_sheet("2. 연도별집계", 1)
        ws3 = wb.create_sheet("3. 월별상세", 2)

        # 스타일 정의
        header_font = Font(bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")

        title_font = Font(bold=True, size=14)
        title_align = Alignment(horizontal="center", vertical="center")

        info_font = Font(bold=True, size=10)
        info_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

        data_align_center = Alignment(horizontal="center", vertical="center")
        data_align_right = Alignment(horizontal="right", vertical="center")

        border_thin = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # ========================================
        # Sheet 1: 입력 정보
        # ========================================
        self._create_sheet1_input_info(ws1, title_font, title_align, info_font, info_fill,
                                        data_align_center, data_align_right, border_thin)

        # ========================================
        # Sheet 2: 연도별 집계표
        # ========================================
        self._create_sheet2_yearly_summary(ws2, title_font, title_align, header_font,
                                             header_fill, header_align, data_align_center,
                                             data_align_right, border_thin)

        # ========================================
        # Sheet 3: 월별 상세표
        # ========================================
        self._create_sheet3_monthly_detail(ws3, title_font, title_align, header_font,
                                             header_fill, header_align, data_align_center,
                                             data_align_right, border_thin)

        # ========================================
        # 인쇄 설정 (모든 시트)
        # ========================================
        for ws in [ws1, ws2, ws3]:
            # 페이지 설정: 가로 1페이지에 맞춤
            ws.page_setup.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0  # 세로는 자동
            ws.page_setup.orientation = 'landscape'  # 가로 방향
            # 여백 설정 (인치 단위)
            ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.75)

        # Sheet 2, 3: 헤더 행 반복 인쇄
        ws2.print_title_rows = '3:3'
        ws3.print_title_rows = '3:3'

        # 파일 저장
        if not output_path:
            # output 폴더 생성 (없으면)
            output_dir = "output"
            if not os.path.exists(output_dir):
                os.makedirs(output_dir)

            # 자산명은 사용자 입력이다 — 경로 구분자·상위참조가 섞이면 output 밖에 쓴다
            safe_name = re.sub(r'[^\w가-힣.-]+', '_', self.asset_input.asset_name).strip('._') or "자산"
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"감가상각스케줄_{safe_name[:60]}_{timestamp}.xlsx"
            output_path = os.path.join(output_dir, filename)

        wb.save(output_path)
        return output_path

    def _calculate_disposal_info(self) -> Optional[Dict[str, Any]]:
        """
        양도로 제거되는 취득원가·누적상각액·장부가액 (시트1 「처분 제거액」)

        분배 산식은 vcore(`disposal.disposal_split`)가 유일한 출처다 — 여기서 다시 계산하면
        같은 파일 안에서 시트1과 시트3(월별)이 어긋난다(감사 G5). 기준월도 vcore와 같은
        **양도월(그 달까지 상각한 뒤)**이다. 직전월이 아니다.

        Returns:
            None (처분 없음) 또는 dict with keys:
                - disposal_date: 처분일자
                - disposal_cost: 처분된 취득원가
                - disposal_accumulated: 처분된 누적상각액
                - disposal_book_value: 처분된 장부가액
                - is_partial: 부분양도 여부
        """
        ai = self.asset_input
        if not ai.disposal_date:
            return None

        disposal_date = datetime.strptime(ai.disposal_date, "%Y-%m-%d")
        row = next((r for r in self.schedule
                    if r['year'] == disposal_date.year and r['month'] == disposal_date.month), None)
        if row is None:
            return None

        # 이벤트 시점 통합 취득원가 = vcore monthly_events의 basis와 같은 정의
        basis = ai.acquisition_cost + (ai.increase_amount or 0)
        cost, accumulated, book_value, is_partial = disposal.disposal_split(
            basis, row['accumulated_dep'], row['book_value'], ai.disposal_amount)

        return {
            'disposal_date': ai.disposal_date,
            'disposal_cost': cost,
            'disposal_accumulated': accumulated,
            'disposal_book_value': book_value,
            'is_partial': is_partial
        }

    def _create_sheet1_input_info(self, ws, title_font, title_align, info_font, info_fill,
                                    data_align_center, data_align_right, border_thin):
        """Sheet 1: 입력 정보"""
        # 제목
        row = 1
        ws.merge_cells(f'A{row}:D{row}')
        cell = ws[f'A{row}']
        cell.value = "감가상각 계산 입력 정보"
        cell.font = title_font
        cell.alignment = title_align

        row += 2

        # 기본 정보
        info_data = [
            ("구분", "항목", "내용", "단위/비고"),
            ("", "", "", ""),
            ("자산 정보", "자산명", self.asset_input.asset_name, ""),
            ("", "자산유형", self.asset_input.asset_type, "유형자산 또는 무형자산"),
            ("", "감가상각방법", self.asset_input.depreciation_method, "정액법 또는 정률법"),
            ("", "", "", ""),
            ("취득 정보", "취득일자", self.asset_input.acquisition_date, "YYYY-MM-DD 형식"),
            ("", "취득원가", self.asset_input.acquisition_cost, "원"),
            ("", "내용연수", self.asset_input.useful_life, "년"),
            ("", "", "", ""),
            ("계산 정보", "잔존가액", 0, "원"),
            ("", "비망가액", 1000, "원"),
        ]

        # 자본적지출 정보 추가 (있는 경우)
        if self.asset_input.increase_date and self.asset_input.increase_amount:
            info_data.extend([
                ("", "", "", ""),
                ("자본적지출 정보", "증가일자", self.asset_input.increase_date, ""),
                ("", "증가금액", self.asset_input.increase_amount, "원"),
            ])

        # 처분 정보 추가 (있는 경우)
        if self.disposal_info:
            disposal_type = "부분양도" if self.disposal_info['is_partial'] else "전체처분"
            label_suffix = "처분부분" if self.disposal_info['is_partial'] else "전체"

            info_data.extend([
                ("", "", "", ""),
                ("처분 정보", "처분일자", self.disposal_info['disposal_date'], ""),
                ("", "처분유형", disposal_type, ""),
                ("", "", "", ""),
                ("처분 제거액", f"취득원가 ({label_suffix})", self.disposal_info['disposal_cost'], "원"),
                ("", f"감가상각누계액 ({label_suffix})", self.disposal_info['disposal_accumulated'], "원"),
                ("", f"장부가액 ({label_suffix})", self.disposal_info['disposal_book_value'], "원"),
            ])

        # 데이터 입력
        for data_row in info_data:
            # 빈 행인지 확인 (모든 값이 빈 문자열)
            is_empty_row = all(v == "" for v in data_row)

            # 빈 행은 건너뛰기 (border 없이 빈 줄만)
            if is_empty_row:
                row += 1
                continue

            for col_idx, value in enumerate(data_row, start=1):
                cell = ws.cell(row=row, column=col_idx, value=value)

                # 첫 번째 행은 헤더
                if row == 3:
                    cell.font = info_font
                    cell.fill = info_fill
                    cell.alignment = data_align_center
                # 구분 열은 볼드
                elif col_idx == 1 and value:
                    cell.font = info_font
                    cell.fill = info_fill
                    cell.alignment = data_align_center
                # 항목 열은 볼드
                elif col_idx == 2 and value:
                    cell.font = info_font
                    cell.alignment = data_align_center
                # 내용 열
                elif col_idx == 3:
                    # 숫자인 경우 천단위 구분 포맷 적용
                    if isinstance(value, int):
                        cell.number_format = '#,##0'
                    cell.alignment = data_align_right
                # 비고 열
                elif col_idx == 4:
                    cell.alignment = data_align_center

                cell.border = border_thin

            row += 1

        # 열 너비 조정
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 30

    def _create_sheet2_yearly_summary(self, ws, title_font, title_align, header_font,
                                        header_fill, header_align, data_align_center,
                                        data_align_right, border_thin):
        """Sheet 2: 연도별 집계표"""
        # 제목
        row = 1
        ws.merge_cells(f'A{row}:F{row}')
        cell = ws[f'A{row}']
        cell.value = f"연도별 감가상각 집계표: {self.asset_input.asset_name}"
        cell.font = title_font
        cell.alignment = title_align

        row += 2

        # 헤더
        headers = ["연도", "상각개월수", "연간 감가상각비", "누적 감가상각비", "기말 장부가액", "비고"]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border_thin

        row += 1

        # 연도별 집계 계산 — 회계연도(결산월 반영) 기준. 달력연도로 묶으면 결산월 ≠ 12에서
        # 월별 시트와 어긋난다(감사 G3). 라벨은 vcore 벡터가 실어준 값을 그대로 쓴다.
        yearly_data = {}
        for monthly in self.schedule:
            year = monthly["fiscal_year"]
            if year not in yearly_data:
                yearly_data[year] = {
                    "months": 0,
                    "total_dep": 0,
                    "ending_accumulated": 0,
                    "ending_book_value": 0
                }

            yearly_data[year]["months"] += 1
            yearly_data[year]["total_dep"] += monthly["monthly_dep"]
            yearly_data[year]["ending_accumulated"] = monthly["accumulated_dep"]
            yearly_data[year]["ending_book_value"] = monthly["book_value"]

        # 데이터 입력
        total_depreciation = 0
        for year in sorted(yearly_data.keys()):
            data = yearly_data[year]
            total_depreciation += data["total_dep"]

            # 비고 결정
            note = ""
            if data["ending_book_value"] == 1000:
                note = "감가상각 완료 (비망가액 1,000원)"
            elif data["months"] < 12:
                if year == min(yearly_data.keys()):
                    note = f"취득년도 ({data['months']}개월)"
                else:
                    note = f"처분/말소년도 ({data['months']}개월)"

            cell = ws.cell(row=row, column=1, value=year)
            cell.alignment = data_align_center
            cell.border = border_thin

            cell = ws.cell(row=row, column=2, value=f"{data['months']}개월")
            cell.alignment = data_align_center
            cell.border = border_thin

            cell = ws.cell(row=row, column=3, value=data['total_dep'])
            cell.number_format = '#,##0'
            cell.alignment = data_align_right
            cell.border = border_thin

            cell = ws.cell(row=row, column=4, value=data['ending_accumulated'])
            cell.number_format = '#,##0'
            cell.alignment = data_align_right
            cell.border = border_thin

            cell = ws.cell(row=row, column=5, value=data['ending_book_value'])
            cell.number_format = '#,##0'
            cell.alignment = data_align_right
            cell.border = border_thin

            cell = ws.cell(row=row, column=6, value=note)
            cell.alignment = data_align_center
            cell.border = border_thin

            row += 1

        # 합계 행
        row += 1
        ws.merge_cells(f'A{row}:B{row}')
        cell = ws[f'A{row}']
        cell.value = "합계"
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border_thin

        cell = ws.cell(row=row, column=3, value=total_depreciation)
        cell.number_format = '#,##0'
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = data_align_right
        cell.border = border_thin

        # 합계 행 나머지 열도 동일한 스타일 적용
        for col in [4, 5, 6]:
            cell = ws.cell(row=row, column=col)
            cell.fill = header_fill
            cell.border = border_thin
            cell.font = header_font  # 흰색 폰트 적용
            cell.alignment = header_align

        # 열 너비 조정
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 30

    def _create_sheet3_monthly_detail(self, ws, title_font, title_align, header_font,
                                        header_fill, header_align, data_align_center,
                                        data_align_right, border_thin):
        """Sheet 3: 월별 상세표"""
        # 제목
        row = 1
        ws.merge_cells(f'A{row}:G{row}')
        cell = ws[f'A{row}']
        cell.value = f"월별 감가상각 상세표: {self.asset_input.asset_name}"
        cell.font = title_font
        cell.alignment = title_align

        row += 2

        # 헤더
        headers = ["연도", "월", "월별 감가상각비", "누적 감가상각비", "장부가액", "계산방법", "비고"]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border_thin

        row += 1

        # 연도별 구분을 위한 스타일 (연도 첫 행에 상단 굵은 테두리)
        border_year_start = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='medium'),  # 상단 굵은 선
            bottom=Side(style='thin')
        )

        # 데이터 입력
        prev_year = None
        for data in self.schedule:
            current_year = data["year"]
            # 연도가 바뀌면 굵은 상단 테두리 사용
            is_year_start = (prev_year is not None and current_year != prev_year)
            current_border = border_year_start if is_year_start else border_thin

            cell = ws.cell(row=row, column=1, value=current_year)
            cell.alignment = data_align_center
            cell.border = current_border

            cell = ws.cell(row=row, column=2, value=f"{data['month']}월")
            cell.alignment = data_align_center
            cell.border = current_border

            cell = ws.cell(row=row, column=3, value=data['monthly_dep'])
            cell.number_format = '#,##0'
            cell.alignment = data_align_right
            cell.border = current_border

            cell = ws.cell(row=row, column=4, value=data['accumulated_dep'])
            cell.number_format = '#,##0'
            cell.alignment = data_align_right
            cell.border = current_border

            cell = ws.cell(row=row, column=5, value=data['book_value'])
            cell.number_format = '#,##0'
            cell.alignment = data_align_right
            cell.border = current_border

            # 계산방법: 가독성 좋게 변환
            formatted_method = _format_calc_method(data['calc_method'])
            cell = ws.cell(row=row, column=6, value=formatted_method)
            cell.alignment = data_align_center
            cell.border = current_border

            cell = ws.cell(row=row, column=7, value=data.get('note', ''))
            cell.alignment = data_align_center
            cell.border = current_border

            prev_year = current_year
            row += 1

        # 열 너비 조정
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 18  # 계산방법 열 너비 조정
        ws.column_dimensions['G'].width = 25


def generate_depreciation_schedule(
    asset_name: str,
    acquisition_date: str,
    acquisition_cost: int,
    useful_life: int,
    asset_type: str,
    depreciation_method: str,
    disposal_date: Optional[str] = None,
    disposal_amount: Optional[int] = None,
    increase_date: Optional[str] = None,
    increase_amount: Optional[int] = None,
    output_path: Optional[str] = None,
    fiscal_year_end_month: int = 12
) -> str:
    """
    감가상각 스케줄을 생성하고 Excel 파일로 출력하는 메인 함수

    Args:
        asset_name: 자산명
        acquisition_date: 취득일자 (YYYY-MM-DD)
        acquisition_cost: 취득원가 (원)
        useful_life: 내용연수 (년)
        asset_type: 자산유형 ("유형자산" or "무형자산")
        depreciation_method: 감가상각방법 ("정액법" or "정률법")
        disposal_date: 처분일자 (YYYY-MM-DD, 선택)
        disposal_amount: 처분금액 (원, 선택)
        increase_date: 증가일자 (YYYY-MM-DD, 선택)
        increase_amount: 증가금액 (원, 선택)
        output_path: 출력 파일 경로 (선택, 없으면 자동생성)

    Returns:
        생성된 Excel 파일 경로

    Example:
        >>> path = generate_depreciation_schedule(
        ...     asset_name="서버장비 A",
        ...     acquisition_date="2020-03-15",
        ...     acquisition_cost=10000000,
        ...     useful_life=5,
        ...     asset_type="유형자산",
        ...     depreciation_method="정액법"
        ... )
        >>> print(f"파일 생성 완료: {path}")
    """
    # 1. 입력 데이터 생성
    asset_input = AssetInput(
        asset_name=asset_name,
        acquisition_date=acquisition_date,
        acquisition_cost=acquisition_cost,
        useful_life=useful_life,
        asset_type=asset_type,
        depreciation_method=depreciation_method,
        disposal_date=disposal_date,
        disposal_amount=disposal_amount,
        increase_date=increase_date,
        increase_amount=increase_amount,
        fiscal_year_end_month=fiscal_year_end_month
    )

    # 2. 스케줄 생성
    generator = MonthlyScheduleGenerator(asset_input)
    schedule = generator.generate_schedule()

    if not schedule:
        raise ValueError("감가상각 스케줄을 생성할 수 없습니다")

    # 3. Excel 파일 생성
    excel_gen = ExcelGenerator(asset_input, schedule)
    output_file = excel_gen.generate_excel(output_path)

    return output_file


if __name__ == "__main__":
    # 테스트 실행
    print("=" * 80)
    print("감가상각 스케줄 자동 생성기 테스트")
    print("=" * 80)

    # 예제 1: 유형자산 정액법
    try:
        output = generate_depreciation_schedule(
            asset_name="서버장비 A",
            acquisition_date="2023-03-15",
            acquisition_cost=10000000,
            useful_life=5,
            asset_type="유형자산",
            depreciation_method="정액법"
        )
        print(f"\n[OK] 예제 1 완료: {output}")
    except Exception as e:
        print(f"\n[ERROR] 예제 1 오류: {e}")

    # 예제 2: 유형자산 정률법
    try:
        output = generate_depreciation_schedule(
            asset_name="차량운반구 B",
            acquisition_date="2022-06-01",
            acquisition_cost=30000000,
            useful_life=5,
            asset_type="유형자산",
            depreciation_method="정률법"
        )
        print(f"[OK] 예제 2 완료: {output}")
    except Exception as e:
        print(f"[ERROR] 예제 2 오류: {e}")

    # 예제 3: 무형자산 정액법
    try:
        output = generate_depreciation_schedule(
            asset_name="특허권 C",
            acquisition_date="2021-01-01",
            acquisition_cost=5000000,
            useful_life=10,
            asset_type="무형자산",
            depreciation_method="정액법"
        )
        print(f"[OK] 예제 3 완료: {output}")
    except Exception as e:
        print(f"[ERROR] 예제 3 오류: {e}")

    print("\n" + "=" * 80)
    print("테스트 완료")
    print("=" * 80)
